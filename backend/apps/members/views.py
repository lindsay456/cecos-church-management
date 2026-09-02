from django_filters.rest_framework import DjangoFilterBackend
from rest_framework import filters, status, viewsets
from rest_framework.decorators import action
from rest_framework.response import Response

from apps.accounts.permissions import ADMIN_ROLES, RoleBasedPermission
from apps.audit.services import audit_create, audit_update
from apps.common.enums import UserRole
from apps.hierarchy.services import entities_in_scope
from apps.members.models import Membre
from apps.members.serializers import (
    MembreCreateSerializer,
    MembreDetailSerializer,
    MembreListSerializer,
    MembreTransferSerializer,
)
from apps.members.services import detect_duplicate_members, transfer_member


class MembreViewSet(viewsets.ModelViewSet):
    serializer_class = MembreListSerializer
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = ["status", "church", "family", "gender", "marital_status"]
    search_fields = ["first_name", "last_name", "email", "phone", "member_number"]
    ordering_fields = ["last_name", "first_name", "membership_date", "created_at"]

    def get_permissions(self):
        if self.action in ("create", "update", "partial_update", "destroy"):
            return [RoleBasedPermission(allowed_roles=ADMIN_ROLES)]
        return [RoleBasedPermission(allowed_roles=ADMIN_ROLES, read_only=True)]

    def get_serializer_class(self):
        if self.action == "create":
            return MembreCreateSerializer
        if self.action == "retrieve":
            return MembreDetailSerializer
        if self.action in ("update", "partial_update"):
            return MembreCreateSerializer
        return MembreListSerializer

    def get_queryset(self):
        qs = Membre.objects.select_related("church", "family").all()
        user = self.request.user
        if user.is_super_admin:
            return qs
        scope = entities_in_scope(user)
        if scope is None:
            return qs.none()
        return qs.filter(church_id__in=scope.values_list("id", flat=True))

    def perform_create(self, serializer):
        kwargs = {}
        if not serializer.validated_data.get("church"):
            user = self.request.user
            if hasattr(user, "entity") and user.entity:
                kwargs["church"] = user.entity
            else:
                from apps.hierarchy.models import EntiteHierarchique
                first_church = EntiteHierarchique.objects.filter(is_active=True).first()
                if first_church:
                    kwargs["church"] = first_church
        if not serializer.validated_data.get("membership_date"):
            from datetime import date
            kwargs["membership_date"] = date.today()
        instance = serializer.save(**kwargs)
        audit_create(instance, user=self.request.user, request=self.request)

    def perform_update(self, serializer):
        old = {f.name: getattr(serializer.instance, f.name) for f in serializer.instance._meta.fields if f.name in ("status", "church_id", "family_id")}
        instance = serializer.save()
        audit_update(instance, user=self.request.user, request=self.request, old_values=old)

    def perform_destroy(self, instance):
        instance.soft_delete()
        audit_update(instance, user=self.request.user, request=self.request, reason="Archivage membre")

    @action(detail=True, methods=["post"])
    def transfer(self, request, pk=None):
        member = self.get_object()
        serializer = MembreTransferSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        from apps.hierarchy.models import EntiteHierarchique
        new_church = EntiteHierarchique.objects.get(id=data["new_church_id"])
        record = transfer_member(
            member, new_church,
            requested_by=request.user,
            reason=data.get("reason", ""),
            notes=data.get("notes", ""),
        )
        return Response({"detail": "Transfert effectué.", "record_id": record.id}, status=status.HTTP_200_OK)

    @action(detail=True, methods=["get"])
    def donations(self, request, pk=None):
        member = self.get_object()
        dons = member.donations.all().order_by("-donation_date")[:50]
        from apps.donations.serializers import DonListSerializer
        return Response(DonListSerializer(dons, many=True).data)

    @action(detail=True, methods=["get"])
    def attendance(self, request, pk=None):
        member = self.get_object()
        presences = member.presences.all().select_related("session").order_by("-session__date")[:50]
        from apps.attendance.serializers import PresenceMembreSerializer
        return Response(PresenceMembreSerializer(presences, many=True).data)

    @action(detail=True, methods=["get"])
    def departments(self, request, pk=None):
        member = self.get_object()
        from apps.departments.models import MembreDepartement
        from apps.departments.serializers import MembreDepartementSerializer
        memberships = MembreDepartement.objects.filter(member=member).select_related("department")
        return Response(MembreDepartementSerializer(memberships, many=True).data)

    @action(detail=True, methods=["get"])
    def receipts(self, request, pk=None):
        member = self.get_object()
        from apps.donations.models import Recu
        from apps.donations.serializers import RecuSerializer
        recus = Recu.objects.filter(donation__member=member).select_related("donation")
        return Response(RecuSerializer(recus, many=True).data)

    @action(detail=True, methods=["get"])
    def transfers(self, request, pk=None):
        member = self.get_object()
        from apps.members.models import HistoriqueAffectationMembre
        from apps.members.serializers import HistoriqueAffectationSerializer
        transfers = HistoriqueAffectationSerializer(member.transfer_history.all()[:20], many=True)
        return Response(transfers.data)

    @action(detail=False, methods=["get"], url_path="export/excel")
    def export_excel(self, request):
        qs = self.filter_queryset(self.get_queryset())
        from apps.reports.generators import generate_report_excel
        headers = ["N°", "Prenom", "Nom", "Genre", "Telephone", "Email", "Statut", "Eglise"]
        rows = []
        for m in qs[:2000]:
            rows.append([
                m.member_number, m.first_name, m.last_name,
                m.get_gender_display() if hasattr(m, 'get_gender_display') else m.gender,
                m.phone or "", m.email or "",
                m.get_status_display(),
                m.church.name if m.church else "",
            ])
        buf = generate_report_excel("Membres", headers, rows)
        from django.http import HttpResponse
        response = HttpResponse(buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = 'attachment; filename="membres.xlsx"'
        return response

    @action(detail=False, methods=["post"], url_path="import/excel")
    def import_excel(self, request):
        file = request.FILES.get("file")
        if not file:
            return Response({"error": "Aucun fichier fourni."}, status=400)
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file)
            ws = wb.active
        except Exception as e:
            return Response({"error": f"Fichier invalide: {e}"}, status=400)

        headers = [cell.value for cell in ws[1]]
        created = 0
        errors = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            data = dict(zip(headers, row))
            first_name = data.get("Prenom") or data.get("prenom") or data.get("first_name")
            last_name = data.get("Nom") or data.get("nom") or data.get("last_name")
            if not first_name or not last_name:
                errors.append(f"Ligne {row_idx}: prenom ou nom manquant")
                continue
            phone = str(data.get("Telephone") or data.get("telephone") or data.get("phone") or "")[:20]
            email = data.get("Email") or data.get("email") or ""
            gender_raw = str(data.get("Genre") or data.get("genre") or data.get("gender") or "M").upper()
            gender = "M" if gender_raw.startswith("M") else "F"

            from apps.members.models import Membre
            church = None
            if hasattr(request.user, 'entity') and request.user.entity:
                church = request.user.entity
            else:
                from apps.hierarchy.models import EntiteHierarchique
                church = EntiteHierarchique.objects.filter(is_active=True).first()

            membre = Membre(
                first_name=first_name,
                last_name=last_name,
                phone=phone,
                email=email,
                gender=gender,
                church=church,
            )
            membre.membership_date = membre.membership_date or __import__('datetime').date.today()
            try:
                membre.full_clean()
                membre.save()
                created += 1
            except Exception as e:
                errors.append(f"Ligne {row_idx}: {str(e)[:80]}")

        return Response({"created": created, "errors": errors[:20], "total_errors": len(errors)})
